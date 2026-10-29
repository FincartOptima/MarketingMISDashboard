# -*- coding: utf-8 -*-
# Writes ld_leads_snapshot.js: hard-coded Workpoint values for the L&D Dashboard tab's clients.
#
# Why it exists: the daily B2C (FIN23) upload only carries leads created this financial year, but
# many L&D clients are older leads, so without this they would show as "not in lead data". Until
# the backend can hold the all-time lead export, this file fills that gap from a one-off all-time
# export. app-ld.js always prefers the live daily B2C record; this is only the fallback, so its
# values go stale only for older leads, and only until it is regenerated.
#
# Only L&D clients are written (not the whole lead export), and each is keyed by a SHA-256 of the
# normalised email rather than the email itself, so the file carries no client emails or names.
#
#   python build_ld_snapshot.py --leads "<all-time FIN export>.xlsx" --ld "<L&D workbook>.xlsx"
#                               [--categories B2C.xlsx]
import argparse
import hashlib
import json
import os
import re
import sys
import unicodedata
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'backend'))
import extract_lib  # noqa: E402

# Keep these three in step with app-ld.js (LD_EMAIL_ALIASES, ldNormKey, ldHashKey).
EMAIL_ALIASES = {
    'inayakpatilvp111@gmail.com': 'vinayakpatilvp111@gmail.com',   # dropped leading "v"
}


def norm_key(v):
    """Email join key: drop every whitespace and invisible/control character, lower-case."""
    if v is None:
        return ''
    text = ''.join(ch for ch in str(v) if not unicodedata.category(ch).startswith('C'))
    key = re.sub(r'\s+', '', text).lower()
    return EMAIL_ALIASES.get(key, key)


def hash_key(key):
    return hashlib.sha256(key.encode('utf-8')).hexdigest()[:16]


def s(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def read_export(path, wanted):
    """{email key: row} for the rows whose email is in `wanted`, from a FIN-style export."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb['Data'] if 'Data' in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [s(x) for x in next(it)]
        col = {h: i for i, h in enumerate(header) if h}
        out = {}
        for r in it:
            key = norm_key(r[col['userId']])
            if key in wanted:
                out[key] = {h: (r[i] if i < len(r) else None) for h, i in col.items()}
        return out
    finally:
        wb.close()


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--leads', required=True, help='all-time FIN lead export (.xlsx, sheet "Data")')
    ap.add_argument('--ld', required=True, help='L&D Financial Planning Tickets Summary workbook (.xlsx)')
    ap.add_argument('--categories', help='B2C export carrying clientCategory (Alpha/Beta/Gamma...)')
    ap.add_argument('--out', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ld_leads_snapshot.js'))
    args = ap.parse_args()

    ld_rows = extract_lib.load_ld_workbook_rows(args.ld)
    # Email column by name, with case and spacing ignored (as app-ld.js reads it).
    email_of = lambda r: next((v for k, v in r.items() if re.sub(r'\s+', ' ', k).strip().lower()
                               in ('email id', 'client mail id') and v), '')
    wanted = {norm_key(email_of(r)) for r in ld_rows} - {''}
    print(f'L&D workbook: {len(ld_rows)} rows, {len(wanted)} distinct client emails')

    print('Reading the all-time lead export (this takes a minute or two)...')
    leads = read_export(args.leads, wanted)
    print(f'  {len(leads)} of those clients found in it')
    cats = {}
    if args.categories:
        for key, r in read_export(args.categories, wanted).items():
            cat = s(r.get('clientCategory'))
            if cat and cat.lower() != 'other':
                cats[key] = cat
        print(f'  {len(cats)} graded (not "Other") in {os.path.basename(args.categories)}')

    rows = {}
    for key in sorted(wanted):
        r = leads.get(key)
        if not r:
            continue
        cat = s(r.get('clientCategory'))
        rows[hash_key(key)] = [
            s(r.get('leadStatus')).upper() or 'BLANK',
            s(r.get('createdDate'))[:10],
            s(r.get('convertedDate'))[:10],
            s(r.get('currentRmName')),
            cats.get(key) or ('' if cat.lower() in ('', 'other') else cat),
            s(r.get('platformName')),
            s(r.get('landingPage')),
        ]
    payload = {
        'generated': datetime.now().strftime('%d %b %Y, %I:%M %p'),
        'leadFile': os.path.basename(args.leads),
        'fields': ['leadStatus', 'createdDate', 'convertedDate', 'currentRmName', 'clientCategory',
                   'platformName', 'landingPage'],
        'rows': rows,
    }
    with open(args.out, 'w', encoding='utf-8') as f:
        f.write('// Generated by build_ld_snapshot.py - do not edit by hand. Workpoint values for L&D\n'
                '// clients, keyed by a hash of the email (see app-ld.js), used where the daily B2C\n'
                '// upload has no record of the client.\n')
        f.write('window.LD_LEAD_SNAPSHOT = ' + json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'Wrote {args.out} ({len(rows)} clients)')


if __name__ == '__main__':
    main()
